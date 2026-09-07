"""Generate demo artifacts from the sample land-record texts.

Produces, for each .txt sample:
  * .pdf   — a clean text-layer PDF (direct extraction path)
  * .docx  — an Office document
  * _scanned.png / _scanned.pdf — an image-only version to demo OCR

Run:  python scripts/generate_samples.py
"""
from __future__ import annotations

import sys
from pathlib import Path

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def _texts() -> list[tuple[str, str]]:
    out = []
    for txt in sorted(SAMPLES.glob("*.txt")):
        out.append((txt.stem, txt.read_text(encoding="utf-8")))
    return out


def make_pdf(path: Path, text: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm

    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    margin = 18 * mm
    y = height - margin
    for line in text.splitlines():
        if y < margin:
            c.showPage()
            y = height - margin
        c.setFont("Helvetica-Bold" if line.isupper() else "Helvetica", 10)
        c.drawString(margin, y, line)
        y -= 4.2 * mm
    c.save()


def make_docx(path: Path, text: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("SovereignAI Sample Document", level=1)
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))


def make_scanned_png(path: Path, text: str) -> None:
    """Render text to a large raster image so Tesseract can read it."""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=28)
    # measure
    probe = Image.new("RGB", (10, 10), "white")
    d = ImageDraw.Draw(probe)
    width = 1240
    line_h = 40
    lines = text.splitlines()
    height = max(200, len(lines) * line_h + 160)
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    y = 60
    for line in lines:
        d.text((80, y), line, fill="black", font=font)
        y += line_h
    img.save(str(path))


def make_scanned_pdf(path: Path, png_path: Path) -> None:
    """Wrap the raster PNG into an image-only PDF (forces OCR on ingestion)."""
    import fitz  # PyMuPDF

    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4 points
    page.insert_image(page.rect, filename=str(png_path))
    doc.save(str(path))
    doc.close()


def main() -> None:
    for stem, text in _texts():
        make_pdf(SAMPLES / f"{stem}.pdf", text)
        make_docx(SAMPLES / f"{stem}.docx", text)
        make_scanned_png(SAMPLES / f"{stem}_scanned.png", text)
        make_scanned_pdf(SAMPLES / f"{stem}_scanned.pdf", SAMPLES / f"{stem}_scanned.png")
        print(f"generated {stem}.pdf / .docx / _scanned.png / _scanned.pdf")


if __name__ == "__main__":
    sys.exit(main())